from io import BytesIO
from pydantic import BaseModel


SUPPORTED_EXTENSIONS = {"txt", "md", "csv", "pdf", "docx"}


class ParsedDocument(BaseModel):
    content: str
    pages: list[str]


def _chunk_text_pages(text: str, page_chars: int = 3000) -> list[str]:
    clean = text.strip()
    if not clean:
        return []
    return [clean[index : index + page_chars] for index in range(0, len(clean), page_chars)]


def parse_document(filename: str, raw: bytes) -> str:
    return parse_document_with_pages(filename, raw).content


def parse_document_with_pages(filename: str, raw: bytes) -> ParsedDocument:
    extension = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Supported files are txt, md, csv, pdf, and docx.")

    if extension == "pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError("PDF parsing requires the pypdf package. Install it with: pip install pypdf") from exc

        reader = PdfReader(BytesIO(raw))
        pages = [page.extract_text() or "" for page in reader.pages]
        clean_pages = [page.strip() for page in pages if page.strip()]
        return ParsedDocument(content="\n\n".join(clean_pages), pages=clean_pages)

    if extension == "docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError(
                "DOCX parsing requires python-docx. If you installed the old docx package, run: "
                "pip uninstall docx -y && pip install python-docx"
            ) from exc

        document = Document(BytesIO(raw))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        content = "\n".join(paragraphs)
        return ParsedDocument(content=content, pages=_chunk_text_pages(content))

    try:
        content = raw.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ValueError("Text-like files must be UTF-8 encoded.") from exc
    return ParsedDocument(content=content, pages=_chunk_text_pages(content))
